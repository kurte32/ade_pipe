"""Read CSV, XLSX and XLS load files into text rows with normalized headers (TDD 8.1).

Every value is kept as text. Excel cell types are respected (dates are rendered as
mm/dd/yyyy HH:MM:SS, numbers without a spurious trailing .0). The real file type is checked by
signature, never by extension alone.
"""

from __future__ import annotations

import csv
import io
import re
import unicodedata
import zipfile
from collections.abc import Iterable, Sequence
from dataclasses import dataclass
from datetime import date, datetime
from pathlib import PurePosixPath

from app.core.config import Settings
from app.core.errors import unprocessable

ALLOWED_EXTENSIONS = {".csv", ".xlsx", ".xls"}
ZIP_MAGIC = b"PK\x03\x04"
OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
DATE_FMT = "%m/%d/%Y %H:%M:%S"


@dataclass
class ParsedFile:
    columns: list[str]
    rows: list[dict[str, str | None]]
    formato: str


def normalize_header(name: object) -> str:
    text = unicodedata.normalize("NFKD", str(name or "").strip())
    text = "".join(ch for ch in text if not unicodedata.combining(ch))
    return text.upper()


def cell_to_text(value: object) -> str | None:
    if value is None:
        return None
    if isinstance(value, datetime):
        return value.strftime(DATE_FMT)
    if isinstance(value, date):
        return value.strftime("%m/%d/%Y")
    if isinstance(value, bool):
        return "1" if value else "0"
    if isinstance(value, float):
        if value.is_integer():
            return str(int(value))
        return repr(value)
    if isinstance(value, int):
        return str(value)
    text = str(value).strip()
    return text if text != "" else None


def _too_many_rows(max_rows: int) -> Exception:
    return unprocessable(f"El archivo supera el límite de {max_rows} filas.")


def _build_rows(
    header: Sequence[object], raw_rows: Iterable[Sequence[object]], max_rows: int
) -> ParsedFile:
    """Convert rows lazily and stop at the first row past ``max_rows`` (audit B-04).

    ``raw_rows`` is consumed as an iterator, so an oversized sheet is rejected after reading
    ``max_rows + 1`` non-empty rows instead of after materialising the whole file.
    """
    columns = [normalize_header(h) for h in header]
    rows: list[dict[str, str | None]] = []
    for raw in raw_rows:
        values = [cell_to_text(v) for v in raw]
        if all(v is None for v in values):
            continue
        if len(rows) >= max_rows:
            raise _too_many_rows(max_rows)
        values += [None] * (len(columns) - len(values))
        rows.append({c: values[i] for i, c in enumerate(columns) if c})
    return ParsedFile(columns=[c for c in columns if c], rows=rows, formato="")


_DIMENSION_RE = re.compile(rb'<dimension ref="[A-Z]+\d+:[A-Z]+(\d+)"')


def _declared_xlsx_rows(data: bytes) -> int | None:
    """Row count from the sheet's `<dimension>` tag, read straight from the zip (audit B-04r).

    Excel and openpyxl write it; a writer that omits it returns None and the streaming count
    applies. The tag sits in the first bytes of the sheet XML, so an oversized file is refused
    before any shared string or cell is parsed.
    """
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            sheets = sorted(n for n in z.namelist() if n.startswith("xl/worksheets/sheet"))
            if not sheets:
                return None
            with z.open(sheets[0]) as fh:
                m = _DIMENSION_RE.search(fh.read(4096))
    except (zipfile.BadZipFile, KeyError, OSError):
        return None
    return int(m.group(1)) if m else None


def _parse_xlsx(data: bytes, max_rows: int) -> ParsedFile:
    declared = _declared_xlsx_rows(data)
    if declared is not None and declared - 1 > max_rows:  # row 1 is the header
        raise _too_many_rows(max_rows)
    if not data.startswith(ZIP_MAGIC):
        raise unprocessable("El archivo no es un XLSX válido (firma incorrecta).")
    from openpyxl import load_workbook

    try:
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception as exc:  # noqa: BLE001
        raise unprocessable("No fue posible leer el XLSX.") from exc
    try:
        ws = wb.worksheets[0]
        it = ws.iter_rows(values_only=True)
        header = list(next(it, []))
        pf = _build_rows(header, it, max_rows)
    finally:
        wb.close()
    pf.formato = "xlsx"
    return pf


def _parse_xls(data: bytes, max_rows: int) -> ParsedFile:
    if not data.startswith(OLE_MAGIC):
        raise unprocessable("El archivo no es un XLS válido (firma incorrecta).")
    import xlrd

    try:
        wb = xlrd.open_workbook(file_contents=data)
    except Exception as exc:  # noqa: BLE001
        raise unprocessable("No fue posible leer el XLS.") from exc
    sh = wb.sheet_by_index(0)
    if sh.nrows == 0:
        return ParsedFile(columns=[], rows=[], formato="xls")

    def conv(cell: xlrd.sheet.Cell) -> object:
        if cell.ctype == xlrd.XL_CELL_DATE:
            return xlrd.xldate_as_datetime(cell.value, wb.datemode)
        if cell.ctype == xlrd.XL_CELL_EMPTY:
            return None
        return cell.value

    header = [c.value for c in sh.row(0)]
    body = ([conv(c) for c in sh.row(i)] for i in range(1, sh.nrows))
    pf = _build_rows(header, body, max_rows)
    pf.formato = "xls"
    return pf


def _parse_csv(data: bytes, max_rows: int) -> ParsedFile:
    if data.startswith(ZIP_MAGIC) or data.startswith(OLE_MAGIC):
        raise unprocessable("El archivo CSV no es texto.")
    try:
        text = data.decode("utf-8-sig")
    except UnicodeDecodeError:
        text = data.decode("latin-1")
    sample = text[:4096]
    try:
        dialect = csv.Sniffer().sniff(sample, delimiters=",;\t|")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect)
    header = next(reader, [])
    pf = _build_rows(list(header), reader, max_rows)
    pf.formato = "csv"
    return pf


def parse_upload(filename: str, data: bytes, settings: Settings) -> ParsedFile:
    ext = PurePosixPath(filename or "").suffix.lower()
    if ext == ".xlsm":
        raise unprocessable("Los archivos con macros (.xlsm) se rechazan.")
    if ext not in ALLOWED_EXTENSIONS:
        raise unprocessable("Formato no admitido. Usa .csv, .xlsx o .xls.")
    if len(data) == 0:
        raise unprocessable("El archivo está vacío.")
    if len(data) > settings.load_max_bytes:
        raise unprocessable(
            f"El archivo supera el límite de {settings.load_max_bytes // (1024 * 1024)} MB."
        )
    if ext == ".xlsx":
        return _parse_xlsx(data, settings.load_max_rows)
    if ext == ".xls":
        return _parse_xls(data, settings.load_max_rows)
    return _parse_csv(data, settings.load_max_rows)
